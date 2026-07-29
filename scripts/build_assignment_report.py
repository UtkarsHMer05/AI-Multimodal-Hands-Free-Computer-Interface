#!/usr/bin/env python3
"""Build the HCI DA1 assignment report and its architecture figure."""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
ASSETS_DIR = ROOT / "assets"
OUTPUT_PATH = ROOT / "AI_Voice_Control_HCI_DA1_Assignment.docx"
ARCHITECTURE_PATH = ASSETS_DIR / "voice_control_architecture.png"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(32, 55, 72)
MUTED = RGBColor(90, 100, 110)
WHITE = RGBColor(255, 255, 255)
LIGHT_FILL = "F4F6F9"
HEADER_FILL = "F4F6F9"
ACCENT_FILL = "E8EEF5"
TABLE_BORDER = "C9D2DC"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


PAPERS = [
    {
        "citation": "Clark et al. (2019)",
        "review": (
            "Systematic mapping of 99 empirical speech-HCI papers. The authors found "
            "strong attention to usability and prototype evaluation, but inconsistent "
            "measurement, limited in-the-wild deployment, and persistent barriers to "
            "building complete speech interfaces. This establishes the project as both "
            "an accessibility and interaction-design problem."
        ),
    },
    {
        "citation": "Deshmukh & Chalmeta (2024)",
        "review": (
            "Systematic review and bibliometric analysis of 125 voice-user-interface "
            "studies. It identifies six research categories and highlights unresolved "
            "issues in accuracy, error recovery, inclusivity, user diversity, context, "
            "and standardized evaluation. These gaps directly motivate constrained "
            "commands, visible feedback, and repeatable metrics."
        ),
    },
    {
        "citation": "Dutsinma et al. (2022)",
        "review": (
            "Reviews voice-assistant usability through ISO 9241-11, separating "
            "effectiveness, efficiency, and satisfaction. The paper shows that voice "
            "studies use varied measures and often omit parts of the ISO framework. "
            "The present methodology therefore measures both objective performance "
            "and subjective experience."
        ),
    },
    {
        "citation": "Harada et al. (2006)",
        "review": (
            "Introduces the Vocal Joystick, which maps vowel quality, loudness, and "
            "pitch to continuous pointer movement. Fitts' law predicted its "
            "speed-accuracy trade-off, and novice tests showed that voice could be a "
            "viable pointing method. It is the closest technical precedent for "
            "voice-based cursor control."
        ),
    },
    {
        "citation": "Harada et al. (2009)",
        "review": (
            "A 2.5-week study with five motor-impaired and four non-impaired users "
            "showed learnability and substantial improvement with the Vocal Joystick. "
            "Motor-impaired participants reached about 70% of previously measured "
            "expert performance. The result supports training effects but also shows "
            "why a simpler discrete-command prototype is appropriate for DA1."
        ),
    },
    {
        "citation": "Ramos et al. (2022)",
        "review": (
            "Develops EMKEY, a low-cost interface combining facial landmarks for "
            "pointer motion with voice commands for mouse/keyboard functions. Tests "
            "with 30 non-disabled and four motor-disabled participants reported "
            "practicality and usability. It validates ordinary webcams, microphones, "
            "and software automation as an affordable assistive approach."
        ),
    },
    {
        "citation": "Corbett & Weber (2016)",
        "review": (
            "Examines a completely hands-free mobile VUI for users with limited hand "
            "dexterity. It identifies command discovery and learning as central "
            "accessibility challenges. The present GUI answers this with a visible "
            "command list, live transcript, explicit status, and deterministic aliases."
        ),
    },
    {
        "citation": "Bekker et al. (1995)",
        "review": (
            "Compares mouse and speech control in a text-annotation system with 24 "
            "participants. Nine used speech more when both modes were available, while "
            "37% preferred access to both. The study shows that speech is valuable as "
            "an alternative rather than a universal replacement for manual input."
        ),
    },
    {
        "citation": "Young & Mihailidis (2010)",
        "review": (
            "Reviews commercial ASR for dysarthric speakers and extends the implications "
            "to older adults. Recognition performance is limited by speech variability "
            "and disorder severity. This prevents the project from claiming universal "
            "accessibility and motivates explicit limitation and target-user testing."
        ),
    },
    {
        "citation": "Berner & Alves (2023)",
        "review": (
            "Scoping review of speech-recognition technology used by people with "
            "disabilities. Only 13 of 78 retrieved articles met the criteria, and many "
            "were non-empirical. The review reports potential gains in participation "
            "and independence but calls for stronger evidence, supporting the proposed "
            "structured user evaluation."
        ),
    },
    {
        "citation": "Cave & Bloch (2023)",
        "review": (
            "Reviews ASR for people living with ALS. Eleven publications linked "
            "recognition performance to dysarthria severity and technology choice, but "
            "functional usability thresholds remain unclear. The project therefore "
            "treats command accuracy and task completion as separate measures."
        ),
    },
    {
        "citation": "Van Schyndel et al. (2014)",
        "review": (
            "Narrative inquiry with three adolescents with physical disabilities and "
            "two parents explains speech-software abandonment through poor fit, effort, "
            "unmet expectations, and easier alternatives. This warns against feature "
            "overload and supports a small vocabulary, fast setup, dry-run mode, and "
            "transparent limitations."
        ),
    },
    {
        "citation": "Yankelovich et al. (1995)",
        "review": (
            "User-informed redesign of SpeechActs argues that speech interfaces must "
            "respect conversational conventions and be designed for speech rather than "
            "copied from graphical interfaces. Error behavior is especially important. "
            "The project uses short commands and direct feedback instead of imitating "
            "open-ended conversation."
        ),
    },
    {
        "citation": "Suhm et al. (2001)",
        "review": (
            "Shows that recognition-error correction can be faster and more accurate "
            "when users can switch modalities. Users learn to avoid ineffective "
            "correction modes as accuracy changes. The prototype retains buttons and "
            "the physical mouse as fallbacks and never forces speech-only recovery."
        ),
    },
    {
        "citation": "Hone & Graham (2000)",
        "review": (
            "Develops the SASSI questionnaire from 214 responses across four speech "
            "applications. Six factors emerged: response accuracy, likeability, "
            "cognitive demand, annoyance, habitability, and speed. These factors provide "
            "a defensible subjective evaluation framework for the proposed study."
        ),
    },
    {
        "citation": "Luger & Sellen (2016)",
        "review": (
            "Interviews with 14 conversational-agent users found a major gap between "
            "expected intelligence and actual system capability. A command-based title, "
            "fixed vocabulary, and explicit supported-command display reduce this gulf "
            "and avoid presenting the prototype as a conversational assistant."
        ),
    },
    {
        "citation": "Cowan et al. (2017)",
        "review": (
            "Qualitative study of infrequent assistant users reports frustration with "
            "incomplete hands-free operation plus concerns about privacy, social "
            "embarrassment, data ownership, and transparency. Local offline recognition "
            "and non-retention of raw audio respond directly to these concerns."
        ),
    },
    {
        "citation": "Schaffer et al. (2015)",
        "review": (
            "Across three experiments, modality efficiency and input error rate strongly "
            "influenced whether users chose touch or speech. The authors developed a "
            "utility-based choice model. This supports keeping manual controls available "
            "and evaluating voice under more than one acoustic condition."
        ),
    },
    {
        "citation": "Oviatt et al. (2004)",
        "review": (
            "Demonstrates that users shift toward multimodal communication as task and "
            "cognitive load rise; multimodal use increased from 18.6% to 77.1% when new "
            "context had to be established. It supports a multimodal safety design in "
            "which speech, buttons, and existing mouse input can coexist."
        ),
    },
    {
        "citation": "Wolters et al. (2009)",
        "review": (
            "Forty-eight users interacted with nine dialogue-system variants. More "
            "options per turn and fewer explicit confirmation sub-dialogues improved "
            "speed without reducing task success in that task. For this prototype, "
            "confirmation is kept concise and visual rather than adding verbal dialogue."
        ),
    },
    {
        "citation": "Wang & Nass (2005)",
        "review": (
            "Two experiments (N=48 and N=96) show that microphone visibility and mobility "
            "can change behavior and attitudes, while output modality had little effect "
            "in those tasks. The finding supports testing with the ordinary built-in "
            "microphone and documenting environmental context."
        ),
    },
    {
        "citation": "Dahlbäck et al. (2007)",
        "review": (
            "A 96-participant experiment found strong preference for system voices with "
            "accents similar to the listener's own, even overriding perceived expertise. "
            "Although this project recognizes rather than synthesizes speech, the study "
            "reinforces the need to test different accents rather than one speaker only."
        ),
    },
    {
        "citation": "Begany et al. (2016)",
        "review": (
            "A 48-participant Wizard-of-Oz comparison of spoken and textual search found "
            "that familiarity, ease of use, speed, trust, comfort, enjoyment, and novelty "
            "shape perception. The evaluation should therefore pair performance logs "
            "with a short subjective questionnaire."
        ),
    },
    {
        "citation": "Limerick et al. (2015)",
        "review": (
            "Two experiments found a reduced sense of agency for voice commands compared "
            "with keyboard input. Immediate transcript and action-result feedback, a "
            "visible enabled/paused state, and predictable one-command-one-action mapping "
            "are design responses intended to strengthen perceived control."
        ),
    },
    {
        "citation": "Sato et al. (2011)",
        "review": (
            "Sasayaki augments a primary auditory browser with secondary contextual "
            "whispers and semantic navigation. Experiments reduced completion time and "
            "increased satisfaction and confidence. The broader lesson is that accessible "
            "voice systems need timely context and focus feedback, not recognition alone."
        ),
    },
    {
        "citation": "Kumar et al. (2012)",
        "review": (
            "Voice Typing displays transcription while the user speaks and supports "
            "rapid touch correction. Compared with traditional dictation, participants "
            "reported lower cognitive demand and made 29% fewer corrections. Live partial "
            "transcription in the prototype follows the same visibility principle."
        ),
    },
    {
        "citation": "Munteanu & Penn (2014)",
        "review": (
            "Argues that speech interfaces are often evaluated through engineering "
            "metrics without enough user-centered design. It highlights both inflated "
            "accuracy expectations and evaluation difficulty. The present methodology "
            "combines recognizer metrics, task behavior, safety checks, and user ratings."
        ),
    },
    {
        "citation": "Koenecke et al. (2020)",
        "review": (
            "Evaluation of five commercial ASR systems on 19.8 hours of speech found "
            "average word error rates of 0.35 for Black speakers and 0.19 for White "
            "speakers. The disparity requires diverse participant sampling and prevents "
            "generalizing a single-speaker smoke test."
        ),
    },
    {
        "citation": "Tatman & Kasten (2017)",
        "review": (
            "Compares Bing Speech and YouTube captions across 39 speakers, four American "
            "English dialects, race, and gender. YouTube showed significant error-rate "
            "differences across dialect and race, while neither system showed a reliable "
            "gender effect. Accent and dialect must be reported in evaluation."
        ),
    },
    {
        "citation": "Chan et al. (2016)",
        "review": (
            "Introduces Listen, Attend and Spell, an end-to-end attention-based neural "
            "recognizer that maps acoustic features directly to characters. It achieved "
            "14.1% word error rate without an external language model and 10.3% with "
            "rescoring. This represents the pretrained deep-learning foundation that "
            "makes application-level ASR practical."
        ),
    },
]


REFERENCES = [
    (
        "L. Clark, P. R. Doyle, D. Garaialde, E. Gilmartin, S. Schlögl, J. Edlund, "
        "M. P. Aylett, J. P. Cabral, C. Munteanu, J. Edwards, and B. R. Cowan, "
        "\"The State of Speech in HCI: Trends, Themes and Challenges,\" "
        "Interacting with Computers, vol. 31, no. 4, pp. 349-371, 2019. ",
        "10.1093/iwc/iwz016",
    ),
    (
        "A. M. Deshmukh and R. Chalmeta, \"User Experience and Usability of Voice "
        "User Interfaces: A Systematic Literature Review,\" Information, vol. 15, "
        "no. 9, art. 579, 2024. ",
        "10.3390/info15090579",
    ),
    (
        "F. L. I. Dutsinma, D. Pal, S. Funilkul, and J. H. Chan, \"A Systematic "
        "Review of Voice Assistant Usability: An ISO 9241-11 Approach,\" SN Computer "
        "Science, vol. 3, art. 267, 2022. ",
        "10.1007/s42979-022-01172-3",
    ),
    (
        "S. Harada, J. A. Landay, J. Malkin, X. Li, and J. A. Bilmes, \"The Vocal "
        "Joystick: Evaluation of Voice-Based Cursor Control Techniques,\" in "
        "Proceedings of ASSETS '06, pp. 197-204, 2006. ",
        "10.1145/1168987.1169021",
    ),
    (
        "S. Harada, J. O. Wobbrock, J. Malkin, J. A. Bilmes, and J. A. Landay, "
        "\"Longitudinal Study of People Learning to Use Continuous Voice-Based "
        "Cursor Control,\" in Proceedings of CHI '09, pp. 347-356, 2009. ",
        "10.1145/1518701.1518757",
    ),
    (
        "P. Ramos, M. Zapata, K. Valencia, V. Vargas, and C. Ramos-Galarza, "
        "\"Low-Cost Human-Machine Interface for Computer Control with Facial "
        "Landmark Detection and Voice Commands,\" Sensors, vol. 22, no. 23, "
        "art. 9279, 2022. ",
        "10.3390/s22239279",
    ),
    (
        "E. Corbett and A. Weber, \"What Can I Say? Addressing User Experience "
        "Challenges of a Mobile Voice User Interface for Accessibility,\" in "
        "Proceedings of MobileHCI '16, pp. 72-82, 2016. ",
        "10.1145/2935334.2935386",
    ),
    (
        "M. M. Bekker, F. L. van Nes, and J. F. Juola, \"A Comparison of Mouse and "
        "Speech Input Control of a Text-Annotation System,\" Behaviour & Information "
        "Technology, vol. 14, no. 1, pp. 14-22, 1995. ",
        "10.1080/01449299508914621",
    ),
    (
        "V. Young and A. Mihailidis, \"Difficulties in Automatic Speech Recognition "
        "of Dysarthric Speakers and Implications for Speech-Based Applications Used "
        "by the Elderly: A Literature Review,\" Assistive Technology, vol. 22, "
        "no. 2, pp. 99-112, 2010. ",
        "10.1080/10400435.2010.483646",
    ),
    (
        "K. Berner and A. N. Alves, \"A Scoping Review of Literature Using Speech "
        "Recognition Technologies by Individuals with Disabilities in Multiple "
        "Contexts,\" Disability and Rehabilitation: Assistive Technology, vol. 18, "
        "no. 7, pp. 1139-1145, 2023. ",
        "10.1080/17483107.2021.1986583",
    ),
    (
        "R. Cave and S. Bloch, \"The Use of Speech Recognition Technology by People "
        "Living with Amyotrophic Lateral Sclerosis: A Scoping Review,\" Disability "
        "and Rehabilitation: Assistive Technology, vol. 18, no. 7, pp. 1043-1055, "
        "2023. ",
        "10.1080/17483107.2021.1974961",
    ),
    (
        "R. Van Schyndel, A. B. Furgoch, T. Previl, and R. Martini, \"The Experience "
        "of Speech Recognition Software Abandonment by Adolescents with Physical "
        "Disabilities,\" Disability and Rehabilitation: Assistive Technology, "
        "vol. 9, no. 6, pp. 513-520, 2014. ",
        "10.3109/17483107.2014.883651",
    ),
    (
        "N. Yankelovich, G.-A. Levow, and M. Marx, \"Designing SpeechActs: Issues in "
        "Speech User Interfaces,\" in Proceedings of CHI '95, pp. 369-376, 1995. ",
        "10.1145/223904.223952",
    ),
    (
        "B. Suhm, B. A. Myers, and A. Waibel, \"Multimodal Error Correction for "
        "Speech User Interfaces,\" ACM Transactions on Computer-Human Interaction, "
        "vol. 8, no. 1, pp. 60-98, 2001. ",
        "10.1145/371127.371166",
    ),
    (
        "K. S. Hone and R. Graham, \"Towards a Tool for the Subjective Assessment "
        "of Speech System Interfaces (SASSI),\" Natural Language Engineering, "
        "vol. 6, nos. 3-4, pp. 287-303, 2000. ",
        "10.1017/S1351324900002497",
    ),
    (
        "E. Luger and A. Sellen, \"Like Having a Really Bad PA: The Gulf Between "
        "User Expectation and Experience of Conversational Agents,\" in Proceedings "
        "of CHI '16, pp. 5286-5297, 2016. ",
        "10.1145/2858036.2858288",
    ),
    (
        "B. R. Cowan, N. Pantidi, D. Coyle, K. Morrissey, P. Clarke, S. Al-Shehri, "
        "D. Earley, and N. Bandeira, \"What Can I Help You With? Infrequent Users' "
        "Experiences of Intelligent Personal Assistants,\" in Proceedings of MobileHCI "
        "'17, pp. 1-12, 2017. ",
        "10.1145/3098279.3098539",
    ),
    (
        "S. Schaffer, R. Schleicher, and S. Möller, \"Modeling Input Modality Choice "
        "in Mobile Graphical and Speech Interfaces,\" International Journal of "
        "Human-Computer Studies, vol. 75, pp. 21-34, 2015. ",
        "10.1016/j.ijhcs.2014.11.004",
    ),
    (
        "S. Oviatt, R. Coulston, and R. Lunsford, \"When Do We Interact Multimodally? "
        "Cognitive Load and Multimodal Communication Patterns,\" in Proceedings of "
        "ICMI '04, pp. 129-136, 2004. ",
        "10.1145/1027933.1027957",
    ),
    (
        "M. Wolters, K. Georgila, J. D. Moore, R. H. Logie, S. E. MacPherson, and "
        "M. Watson, \"Reducing Working Memory Load in Spoken Dialogue Systems,\" "
        "Interacting with Computers, vol. 21, no. 4, pp. 276-287, 2009. ",
        "10.1016/j.intcom.2009.05.009",
    ),
    (
        "Q. Wang and C. Nass, \"Less Visible and Wireless: Two Experiments on the "
        "Effects of Microphone Type on Users' Performance and Perception,\" in "
        "Proceedings of CHI '05, pp. 809-818, 2005. ",
        "10.1145/1054972.1055086",
    ),
    (
        "N. Dahlbäck, Q. Wang, C. Nass, and J. Alwin, \"Similarity Is More Important "
        "Than Expertise: Accent Effects in Speech Interfaces,\" in Proceedings of "
        "CHI '07, pp. 1553-1556, 2007. ",
        "10.1145/1240624.1240859",
    ),
    (
        "G. M. Begany, N. Sa, and X. Yuan, \"Factors Affecting User Perception of a "
        "Spoken Language vs. Textual Search Interface: A Content Analysis,\" "
        "Interacting with Computers, vol. 28, no. 2, pp. 170-180, 2016. ",
        "10.1093/iwc/iwv029",
    ),
    (
        "H. Limerick, J. W. Moore, and D. Coyle, \"Empirical Evidence for a "
        "Diminished Sense of Agency in Speech Interfaces,\" in Proceedings of CHI "
        "'15, pp. 3967-3970, 2015. ",
        "10.1145/2702123.2702379",
    ),
    (
        "D. Sato, S. Zhu, M. Kobayashi, H. Takagi, and C. Asakawa, \"Sasayaki: "
        "Augmented Voice Web Browsing Experience,\" in Proceedings of CHI '11, "
        "pp. 2769-2778, 2011. ",
        "10.1145/1978942.1979353",
    ),
    (
        "A. Kumar, T. Paek, and B. Lee, \"Voice Typing: A New Speech Interaction "
        "Model for Dictation on Touchscreen Devices,\" in Proceedings of CHI '12, "
        "pp. 2277-2286, 2012. ",
        "10.1145/2207676.2208386",
    ),
    (
        "C. Munteanu and G. Penn, \"Speech-Based Interaction: Myths, Challenges, "
        "and Opportunities,\" in CHI '14 Extended Abstracts, pp. 1035-1036, 2014. ",
        "10.1145/2559206.2567826",
    ),
    (
        "A. Koenecke, A. Nam, E. Lake, J. Nudell, M. Quartey, Z. Mengesha, "
        "C. Toups, J. R. Rickford, D. Jurafsky, and S. Goel, \"Racial Disparities "
        "in Automated Speech Recognition,\" Proceedings of the National Academy "
        "of Sciences, vol. 117, no. 14, pp. 7684-7689, 2020. ",
        "10.1073/pnas.1915768117",
    ),
    (
        "R. Tatman and C. Kasten, \"Effects of Talker Dialect, Gender & Race on "
        "Accuracy of Bing Speech and YouTube Automatic Captions,\" in Interspeech "
        "2017, pp. 934-938, 2017. ",
        "10.21437/Interspeech.2017-1746",
    ),
    (
        "W. Chan, N. Jaitly, Q. V. Le, and O. Vinyals, \"Listen, Attend and Spell: "
        "A Neural Network for Large Vocabulary Conversational Speech Recognition,\" "
        "in IEEE ICASSP 2016, pp. 4960-4964, 2016. ",
        "10.1109/ICASSP.2016.7472621",
    ),
]


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
        if bold
        else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size=size, index=0)
        except OSError:
            continue
    return ImageFont.load_default()


def create_architecture_diagram(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    width, height = 1800, 720
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = font(40, bold=True)
    box_title_font = font(25, bold=True)
    body_font = font(20)
    arrow_font = font(34, bold=True)

    draw.text(
        (width // 2, 42),
        "Offline voice-command processing pipeline",
        font=title_font,
        fill=(32, 55, 72),
        anchor="ma",
    )

    boxes = [
        ("1  Microphone", "16 kHz mono\naudio stream"),
        ("2  Vosk ASR", "Pretrained offline\nspeech model"),
        ("3  Grammar", "14 supported\ncommand phrases"),
        ("4  Parser", "Normalize and map\nexact aliases"),
        ("5  Safety gate", "Paused / enabled\nstate check"),
        ("6  Action + UX", "PyAutoGUI action\nfeedback and CSV log"),
    ]
    box_w, box_h = 245, 245
    gap = 44
    total = len(boxes) * box_w + (len(boxes) - 1) * gap
    x0 = (width - total) // 2
    y0 = 185

    for index, (heading, detail) in enumerate(boxes):
        x = x0 + index * (box_w + gap)
        fill = (232, 238, 245) if index != 4 else (244, 246, 249)
        draw.rounded_rectangle(
            (x, y0, x + box_w, y0 + box_h),
            radius=18,
            fill=fill,
            outline=(46, 116, 181),
            width=4,
        )
        draw.text(
            (x + box_w // 2, y0 + 55),
            heading,
            font=box_title_font,
            fill=(31, 77, 120),
            anchor="mm",
        )
        draw.multiline_text(
            (x + box_w // 2, y0 + 145),
            detail,
            font=body_font,
            fill=(45, 50, 55),
            anchor="mm",
            align="center",
            spacing=8,
        )
        if index < len(boxes) - 1:
            draw.text(
                (x + box_w + gap // 2, y0 + box_h // 2),
                "→",
                font=arrow_font,
                fill=(46, 116, 181),
                anchor="mm",
            )

    draw.rounded_rectangle(
        (300, 520, 1500, 652),
        radius=16,
        fill=(248, 249, 251),
        outline=(201, 210, 220),
        width=3,
    )
    draw.text(
        (900, 555),
        "Safety invariant",
        font=box_title_font,
        fill=(31, 77, 120),
        anchor="ma",
    )
    draw.text(
        (900, 610),
        "Evaluation mode and paused mode recognize commands but never execute computer actions.",
        font=body_font,
        fill=(45, 50, 55),
        anchor="ma",
    )
    image.save(path, dpi=(180, 180))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must total {CONTENT_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(
    run,
    *,
    name: str = "Calibri",
    size: float | None = None,
    color: RGBColor | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_repeatable_cell_text(cell, size=9.2, bold=False, color=None, align=None) -> None:
    for paragraph in cell.paragraphs:
        if align is not None:
            paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.08
        for run in paragraph.runs:
            set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_hyperlink(paragraph, text: str, url: str):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(r_fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    r_pr.append(size)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_field(paragraph, instruction: str, display: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def create_numbering(doc: Document, *, fmt: str, text: str, left=540, hanging=280) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl.append(lvl_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    lvl.append(suffix)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    p_pr.append(ind)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])


def add_bullets(doc: Document, items: list[str]) -> None:
    num_id = create_numbering(doc, fmt="bullet", text="•", left=540, hanging=280)
    for item in items:
        paragraph = doc.add_paragraph()
        apply_numbering(paragraph, num_id)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.208
        paragraph.add_run(item)


def add_steps(doc: Document, items: list[str]) -> None:
    num_id = create_numbering(doc, fmt="decimal", text="%1. ", left=540, hanging=280)
    for item in items:
        paragraph = doc.add_paragraph()
        apply_numbering(paragraph, num_id)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.208
        paragraph.add_run(item)


def add_callout(doc: Document, label: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_FILL)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    lead = paragraph.add_run(f"{label}: ")
    set_run_font(lead, bold=True, color=DARK_BLUE)
    paragraph.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    set_run_font(run, size=9.5, color=MUTED, italic=True)


def add_section_intro(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(text)
    set_run_font(run, size=11, color=INK, bold=True)


def add_heading(doc: Document, text: str, level: int = 1):
    return doc.add_heading(text, level=level)


def add_text(doc: Document, text: str, *, bold_lead: str | None = None):
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        paragraph.add_run(text[len(bold_lead) :])
    else:
        paragraph.add_run(text)
    return paragraph


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    heading_tokens = {
        1: (16, BLUE, 18, 10),
        2: (13, BLUE, 12, 6),
        3: (12, DARK_BLUE, 8, 4),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    if "Table Text" not in styles:
        table_text = styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_text = styles["Table Text"]
    table_text.font.name = "Calibri"
    table_text._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    table_text._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    table_text.font.size = Pt(9.2)
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.space_after = Pt(2)
    table_text.paragraph_format.line_spacing = 1.08

    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def set_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("HCI DA1  |  Voice-Controlled Computer Interface")
    set_run_font(run, size=9, color=MUTED, bold=True)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    add_field(paragraph, "PAGE", "1")


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(70)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run("HUMAN-COMPUTER INTERACTION  |  DA1")
    set_run_font(run, size=11, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run("AI-Based Voice-Controlled\nComputer Interface")
    set_run_font(run, size=28, color=INK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(34)
    run = subtitle.add_run("Hands-Free Cursor and Basic Action Control")
    set_run_font(run, size=16, color=DARK_BLUE)

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    line.paragraph_format.space_after = Pt(40)
    run = line.add_run("Topic finalization • 30-paper literature review • Hardware • Software • Methodology")
    set_run_font(run, size=10.5, color=MUTED, italic=True)

    for label, value in [
        ("Submitted by", "________________________________________"),
        ("Roll number", "________________________________________"),
        ("Course / section", "________________________________________"),
        ("Submission date", "29 July 2026"),
    ]:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(7)
        lead = paragraph.add_run(f"{label}: ")
        set_run_font(lead, size=11, color=INK, bold=True)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=11, color=INK)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(36)
    run = note.add_run("Deliverable: Word report (no presentation deck required)")
    set_run_font(run, size=9.5, color=MUTED)
    doc.add_page_break()


def add_requirements_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = [900, 4200, 900, 3360]
    set_table_geometry(table, widths)
    headers = ["Item", "Assessment requirement", "Marks", "Evidence in report"]
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        set_cell_shading(cell, HEADER_FILL)
        set_repeatable_cell_text(
            cell, size=9.2, bold=True, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER
        )
    repeat_table_header(table.rows[0])
    rows = [
        ("1", "Review of related literature (25-30 papers)", "4", "Section 3: 30 verified papers and synthesis"),
        ("2", "Required hardware components", "2", "Section 5: mandatory and optional hardware"),
        ("3", "Required software", "2", "Section 6: platform, libraries, and model"),
        ("4", "Finalized methodology", "2", "Section 7: pipeline, experiment, metrics, and ethics"),
    ]
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].text = value
            align = WD_ALIGN_PARAGRAPH.CENTER if index in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
            set_repeatable_cell_text(cells[index], align=align)
    set_table_geometry(table, widths)


def add_contents(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "Contents", 1)
    add_section_intro(
        doc,
        "Report roadmap: the document follows the DA1 criteria first, then records "
        "implementation and verification evidence."
    )
    add_steps(
        doc,
        [
            "Topic Finalization - selected topic, rationale, problem, aim, objectives, and research questions.",
            "Review Method - search concepts, inclusion criteria, and evidence boundaries.",
            "Review of Related Literature - thematic synthesis, 30-paper matrix, and research gap.",
            "Proposed System and Current Implementation - scope, architecture, safety, and dataset decision.",
            "Required Hardware Components - mandatory, fallback, and optional devices.",
            "Required Software - fixed software stack, project modules, and setup sequence.",
            "Finalized Methodology - development method, human evaluation, measures, analysis, and ethics.",
            "Implementation Verification and Expected Outcomes - completed checks and limitations.",
            "Conclusion and References - final decision plus 30 DOI-linked scholarly references.",
        ],
    )
    doc.add_page_break()


def add_literature_matrix(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    widths = [520, 2320, 6520]
    set_table_geometry(table, widths)
    headers = ["No.", "Study", "Critical review and relevance to the proposed system"]
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        set_cell_shading(cell, HEADER_FILL)
        set_repeatable_cell_text(
            cell, size=9.2, bold=True, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER
        )
    repeat_table_header(table.rows[0])

    for index, paper in enumerate(PAPERS, start=1):
        cells = table.add_row().cells
        cells[0].text = str(index)
        cells[1].text = f"[{index}] {paper['citation']}"
        cells[2].text = paper["review"]
        set_repeatable_cell_text(cells[0], align=WD_ALIGN_PARAGRAPH.CENTER)
        set_repeatable_cell_text(cells[1], bold=True, color=INK)
        set_repeatable_cell_text(cells[2])
    set_table_geometry(table, widths)


def add_hardware_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = [1800, 1800, 3100, 2660]
    set_table_geometry(table, widths)
    headers = ["Component", "Status", "Minimum / example", "Purpose"]
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        set_cell_shading(cell, HEADER_FILL)
        set_repeatable_cell_text(
            cell, bold=True, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER
        )
    repeat_table_header(table.rows[0])
    rows = [
        ("Laptop / desktop", "Required", "64-bit CPU; 4 GB RAM minimum; 8 GB recommended", "Runs the GUI, offline ASR model, and automation."),
        ("Microphone", "Required", "Built-in or USB; mono input; 16 kHz supported", "Captures spoken commands."),
        ("Display", "Required", "Any standard monitor", "Shows transcript, state, feedback, and evaluation prompts."),
        ("Mouse / trackpad", "Safety fallback", "Existing device", "Allows immediate manual recovery and fail-safe movement."),
        ("Speakers / headphones", "Optional", "Any audio output", "Useful only if spoken feedback is added later."),
        ("External USB microphone", "Optional", "Noise-reducing headset or desktop mic", "May improve recognition in noisy rooms."),
        ("GPU / sensors", "Not required", "No dedicated GPU, bend sensor, eye tracker, or wearable", "Keeps the prototype low-cost and reproducible."),
    ]
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].text = value
            align = WD_ALIGN_PARAGRAPH.CENTER if index == 1 else WD_ALIGN_PARAGRAPH.LEFT
            set_repeatable_cell_text(cells[index], align=align)
    set_table_geometry(table, widths)


def add_software_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = [1800, 1350, 2600, 3610]
    set_table_geometry(table, widths)
    headers = ["Software", "Version", "Role", "Reason selected"]
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        set_cell_shading(cell, HEADER_FILL)
        set_repeatable_cell_text(
            cell, bold=True, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER
        )
    repeat_table_header(table.rows[0])
    rows = [
        ("Python", "3.10+", "Application language", "Readable, portable, and well supported for HCI prototypes."),
        ("Vosk", "0.3.44", "Offline automatic speech recognition", "Works locally and accepts a constrained runtime grammar."),
        ("English Vosk model", "small-en-us-0.15", "Pretrained acoustic/language model", "Avoids collection and training of a custom dataset."),
        ("sounddevice", "0.5.5", "16 kHz microphone streaming", "Provides cross-platform access to PortAudio devices."),
        ("PyAutoGUI", "0.9.54", "Mouse, click, scroll, and shortcuts", "Simple OS automation with an emergency corner fail-safe."),
        ("Tkinter", "Python standard library", "Desktop user interface", "No separate GUI framework or web server is required."),
        ("CSV / threading / webbrowser", "Python standard library", "Logging, background recognition, browser action", "Reduces dependency count and keeps data readable."),
        ("unittest", "Python standard library", "Automated verification", "Supports repeatable tests without extra tooling."),
    ]
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].text = value
            align = WD_ALIGN_PARAGRAPH.CENTER if index == 1 else WD_ALIGN_PARAGRAPH.LEFT
            set_repeatable_cell_text(cells[index], align=align)
    set_table_geometry(table, widths)


def add_test_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = [2150, 3250, 1860, 2100]
    set_table_geometry(table, widths)
    headers = ["Verification layer", "What was checked", "Evidence", "Result"]
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        set_cell_shading(cell, HEADER_FILL)
        set_repeatable_cell_text(
            cell, bold=True, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER
        )
    repeat_table_header(table.rows[0])
    rows = [
        ("Unit tests", "Parsing, aliases, unsafe phrase rejection, pause/resume, dry-run, evaluation statistics, CSV export", "13 automated tests", "Passed"),
        ("Dependency check", "Installed package consistency on Apple Silicon / Python 3.12", "pip check", "No broken requirements"),
        ("Model integration", "Local model loading and constrained Vosk grammar", "Recognizer construction", "Passed"),
        ("Audio-device check", "Available input devices on development laptop", "Three microphone-capable devices detected", "Passed"),
        ("Synthetic command smoke test", "Primary phrase for all 14 commands, synthesized voice, quiet controlled audio", "14/14 recognized", "Passed; not a human-user result"),
        ("GUI launch", "Dry-run application startup", "Window process launched successfully", "Passed"),
    ]
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].text = value
            align = WD_ALIGN_PARAGRAPH.CENTER if index == 3 else WD_ALIGN_PARAGRAPH.LEFT
            set_repeatable_cell_text(cells[index], align=align)
    set_table_geometry(table, widths)


def add_references(doc: Document) -> None:
    num_id = create_numbering(doc, fmt="decimal", text="[%1]", left=540, hanging=540)
    for citation, doi in REFERENCES:
        paragraph = doc.add_paragraph()
        apply_numbering(paragraph, num_id)
        paragraph.paragraph_format.left_indent = Inches(0.38)
        paragraph.paragraph_format.first_line_indent = Inches(-0.38)
        paragraph.paragraph_format.space_after = Pt(7)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.keep_together = True
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(citation)
        set_run_font(run, size=10)
        url = f"https://doi.org/{doi}"
        add_hyperlink(paragraph, url, url)


def build_report() -> Path:
    create_architecture_diagram(ARCHITECTURE_PATH)
    doc = Document()
    style_document(doc)
    set_header_footer(doc)
    add_cover(doc)

    add_heading(doc, "Executive Summary", 1)
    add_callout(
        doc,
        "Final topic",
        "AI-Based Voice-Controlled Computer Interface for Hands-Free Cursor and Basic Action Control.",
    )
    add_text(
        doc,
        "This report finalizes a focused version of the proposed topic “Voice based control "
        "of any action.” The selected project is a command-based desktop interface that "
        "recognizes a small vocabulary locally and maps each valid phrase to one predictable "
        "computer action. The implementation uses a pretrained Vosk speech model, so no "
        "custom training dataset, cloud API, specialized sensor, or GPU is required."
    )
    add_text(
        doc,
        "The project is technically complete as a working prototype: it supports cursor "
        "movement, clicking, scrolling, browser and tab actions; starts in a safe paused "
        "state; includes a dry-run mode; displays recognition and action feedback; writes CSV "
        "logs; and provides a guided accuracy/response-time evaluation. The literature review "
        "covers exactly 30 DOI-verified papers spanning speech HCI, assistive access, cursor "
        "control, error recovery, usability measurement, multimodality, and ASR fairness."
    )
    add_heading(doc, "Assessment Requirements Addressed", 2)
    add_requirements_table(doc)
    add_caption(doc, "Table 1. Direct mapping from DA1 marking criteria to report evidence.")
    add_callout(
        doc,
        "Scope boundary",
        "The system is a fixed-command accessibility prototype, not a conversational assistant "
        "and not a clinical device. Human-user testing is proposed but no unperformed study is "
        "reported as a completed result.",
    )
    add_contents(doc)

    add_heading(doc, "1. Topic Finalization", 1)
    add_section_intro(
        doc,
        "Selected topic: AI-Based Voice-Controlled Computer Interface for Hands-Free Cursor "
        "and Basic Action Control."
    )
    add_heading(doc, "1.1 Rationale for Selection", 2)
    add_text(
        doc,
        "Among the proposed topics, this option provides the best balance of implementation "
        "simplicity, visible HCI value, and literature depth. A laptop already contains the "
        "necessary microphone, display, processor, mouse/trackpad fallback, and operating "
        "system. A pretrained recognizer removes model-training effort, while a constrained "
        "grammar keeps the interaction testable and reduces accidental activation."
    )
    add_bullets(
        doc,
        [
            "No EEG headset, eye tracker, bend sensor, wearable, robot, AR/VR headset, or custom control panel is required.",
            "No image collection, annotation, gesture dataset, or deep-learning training pipeline is required.",
            "The result is immediately demonstrable: a spoken phrase causes a visible cursor or system action.",
            "The topic has mature HCI, accessibility, usability, speech-recognition, and fairness literature.",
            "The prototype can be evaluated quantitatively through command accuracy and response time.",
        ],
    )
    add_heading(doc, "1.2 Problem Statement", 2)
    add_text(
        doc,
        "Conventional mouse and keyboard input can be difficult or unavailable for people "
        "with upper-limb motor limitations and can also be inconvenient when the hands are "
        "occupied. Existing general-purpose voice assistants are often open-ended, cloud "
        "dependent, difficult to discover, and inconsistent in how they map language to "
        "desktop actions. The problem is to provide a low-cost, offline, predictable, and "
        "measurable voice interface for a small set of common computer operations."
    )
    add_heading(doc, "1.3 Aim and Objectives", 2)
    add_text(
        doc,
        "Aim: To design, implement, and evaluate an offline voice-command interface that "
        "supports hands-free cursor and basic computer control using ordinary laptop hardware."
    )
    add_bullets(
        doc,
        [
            "Recognize a fixed vocabulary of 14 supported commands with a pretrained offline model.",
            "Map each recognized command deterministically to one cursor, click, scroll, browser, tab, or control-state action.",
            "Provide visible partial transcription, final feedback, command discoverability, and session logging.",
            "Prevent unintended actions through a paused startup state, constrained grammar, dry-run mode, and PyAutoGUI fail-safe.",
            "Measure command accuracy and response time through a guided evaluation that never executes actions.",
        ],
    )
    add_heading(doc, "1.4 Research Questions", 2)
    add_steps(
        doc,
        [
            "How accurately does a constrained offline recognizer identify the supported commands in quiet and noisy conditions?",
            "How quickly can users complete basic cursor and desktop-control tasks using discrete voice commands?",
            "How do users rate response accuracy, speed, cognitive demand, annoyance, likeability, and habitability?",
            "Which errors, accents, acoustic conditions, and interaction states create the greatest accessibility risks?",
        ],
    )

    add_heading(doc, "2. Review Method", 1)
    add_text(
        doc,
        "The literature set was constructed as a focused related-literature review rather "
        "than a claim of a new systematic review. Searches covered DOI-indexed publications "
        "from ACM, IEEE, Oxford University Press, Elsevier, Springer, Taylor & Francis, MDPI, "
        "Cambridge University Press, ISCA, and PNAS. Backward searching from major speech-HCI "
        "reviews was used to identify foundational studies."
    )
    add_heading(doc, "2.1 Search Concepts", 2)
    add_bullets(
        doc,
        [
            "“voice cursor control” OR “speech-based cursor control”",
            "“voice user interface” AND usability OR accessibility",
            "automatic speech recognition AND disability OR motor impairment",
            "speech interface AND error correction OR cognitive load OR agency",
            "automatic speech recognition AND accent OR dialect OR racial disparity",
        ],
    )
    add_heading(doc, "2.2 Inclusion and Exclusion", 2)
    add_text(
        doc,
        "Included papers were peer-reviewed journal or conference works with a DOI, directly "
        "relevant to speech interaction, assistive access, cursor/desktop control, VUI design, "
        "evaluation, or ASR performance. Systematic and scoping reviews were retained to map "
        "the evidence base. Papers without a verifiable DOI, purely unrelated speech "
        "applications, and non-scholarly product pages were excluded. The final set contains "
        "30 papers published from 1995 to 2024."
    )

    add_heading(doc, "3. Review of Related Literature", 1)
    add_section_intro(
        doc,
        "The evidence converges on five design requirements: accessible alternatives, "
        "predictable commands, recoverable errors, measurable usability, and inclusive testing."
    )
    add_heading(doc, "3.1 State of Speech HCI and Evaluation", 2)
    add_text(
        doc,
        "The two broadest reviews show that literature volume is not a limitation. Clark et "
        "al. map 99 empirical speech-HCI papers and identify nine themes, including assistive "
        "technology and accessibility [1]. Deshmukh and Chalmeta analyze 125 VUI papers and "
        "highlight recurring gaps in error recovery, inclusivity, user variety, and evaluation "
        "[2]. Dutsinma et al. further show that effectiveness, efficiency, and satisfaction are "
        "not measured consistently [3]. Together they justify a small, fully implemented system "
        "with explicit metrics instead of a broad conceptual assistant."
    )
    add_heading(doc, "3.2 Assistive Cursor and Computer Control", 2)
    add_text(
        doc,
        "The Vocal Joystick work is the direct technical foundation. Its continuous mapping "
        "from vocal parameters to pointer direction demonstrated feasibility and learnability "
        "[4], [5]. EMKEY shows that low-cost computer control can combine conventional cameras "
        "and voice commands [6]. Mobile accessibility research emphasizes that fully hands-free "
        "operation matters for limited dexterity but command discovery remains difficult [7]. "
        "Earlier comparison work also shows that many users value speech as one available "
        "modality rather than an exclusive replacement for the mouse [8]."
    )
    add_heading(doc, "3.3 Accessibility Limits and Technology Abandonment", 2)
    add_text(
        doc,
        "Speech recognition does not automatically guarantee accessibility. Dysarthric speech "
        "and age-related voice changes can reduce recognition performance [9]. Reviews of "
        "disability contexts and ALS show potential for independence but limited empirical "
        "evidence and an unclear relationship between word error rate and real usability "
        "[10], [11]. Abandonment research shows that users reject systems that demand too much "
        "effort or fit poorly with their tasks [12]. Therefore this project does not claim "
        "universal accessibility and preserves manual fallback."
    )
    add_heading(doc, "3.4 Discoverability, Errors, Feedback, and Agency", 2)
    add_text(
        doc,
        "SpeechActs established that VUIs must be designed around speech conventions and error "
        "conditions rather than copied from graphical interfaces [13]. Multimodal correction "
        "is more effective than repeatedly speaking when recognition fails [14]. SASSI "
        "provides six subjective factors for evaluation [15]. Studies of assistants find "
        "expectation gaps, incomplete hands-free interaction, privacy concerns, and social "
        "discomfort [16], [17]. Reduced agency under voice control [24] and benefits of visible "
        "real-time transcription [26] support the prototype's transcript, state indicator, "
        "one-command-one-action mapping, and concise result feedback."
    )
    add_heading(doc, "3.5 Context, Multimodality, and Inclusive ASR", 2)
    add_text(
        doc,
        "Users change modality according to error rate, task efficiency, cognitive load, and "
        "context [18], [19]. Microphone form and accent can change behavior and perceived "
        "expertise [21], [22], while spoken-interface perception also depends on trust, speed, "
        "comfort, and familiarity [23]. More importantly, ASR accuracy can differ substantially "
        "across race and dialect [28], [29]. These results require diverse evaluation and "
        "careful interpretation of any accuracy number. Modern end-to-end neural recognition "
        "architectures explain why a pretrained application-level solution is feasible [30]."
    )
    add_heading(doc, "3.6 Critical Review Matrix (30 Papers)", 2)
    add_literature_matrix(doc)
    add_caption(doc, "Table 2. Critical review matrix of the 30 DOI-verified related papers.")
    add_heading(doc, "3.7 Research Gap", 2)
    add_callout(
        doc,
        "Identified gap",
        "Prior research establishes continuous vocal pointing, general voice assistants, "
        "multimodal accessibility, and ASR evaluation, but a compact student-ready prototype "
        "can contribute by combining offline fixed-command recognition, explicit safety state, "
        "visible feedback, automatic CSV logging, and a guided non-executing evaluation in one "
        "reproducible desktop application.",
    )
    add_text(
        doc,
        "The project does not claim a new ASR algorithm. Its contribution is an HCI design and "
        "evaluation integration: it applies established findings about discoverability, error "
        "visibility, manual fallback, privacy, and user diversity to a working low-cost system."
    )

    add_heading(doc, "4. Proposed System and Current Implementation", 1)
    add_heading(doc, "4.1 Functional Scope", 2)
    add_text(
        doc,
        "The application recognizes 14 command classes. Twelve perform or describe computer "
        "actions and two control the safety state. The primary phrases are: move left, move "
        "right, move up, move down, click, double click, right click, scroll up, scroll down, "
        "open browser, new tab, close tab, pause control, and resume control. A few explicit "
        "aliases such as “go left” and “enable control” are accepted; arbitrary language is not."
    )
    add_heading(doc, "4.2 Architecture", 2)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    inline = run.add_picture(str(ARCHITECTURE_PATH), width=Inches(6.35))
    doc_pr = inline._inline.docPr
    doc_pr.set("descr", "Six-stage offline voice-command processing pipeline with a safety invariant.")
    add_caption(doc, "Figure 1. Implemented offline voice-command architecture.")
    add_text(
        doc,
        "Microphone audio is streamed at 16 kHz to Vosk. A constrained grammar produces only "
        "supported phrases or an unknown token. Text is normalized and exactly matched to a "
        "command object. The action executor checks the paused/enabled state before using "
        "PyAutoGUI or the browser module. Every result is shown in the GUI and regular "
        "interaction is written to a timestamped CSV file."
    )
    add_heading(doc, "4.3 Safety and Accessibility Features", 2)
    add_bullets(
        doc,
        [
            "Control starts paused; movement and action commands are ignored until explicitly enabled.",
            "Dry-run mode performs recognition, feedback, and logging without any OS action.",
            "Guided evaluation mode records recognition results but never executes commands.",
            "PyAutoGUI's corner fail-safe remains enabled for emergency interruption.",
            "The supported-command list, live partial transcript, final transcript, and action result remain visible.",
            "Recognition runs locally after model download; raw microphone audio is not saved.",
        ],
    )
    add_heading(doc, "4.4 Dataset and Model Decision", 2)
    add_callout(
        doc,
        "No custom training dataset required",
        "The system uses the pretrained Vosk small English model and constrains decoding to "
        "the project's grammar. Training a new network would add data collection, annotation, "
        "compute, bias, and validation work without improving the DA1 objective.",
    )
    add_text(
        doc,
        "The installed model occupies approximately 68 MB. Evaluation data is created by the "
        "application itself as CSV trials containing target command, recognized text, matched "
        "command, correctness, and response time. Raw audio is deliberately not retained. If "
        "future work requires adaptation for Indian English or dysarthric speech, a consented, "
        "balanced corpus and separate ethics/data-management plan would be necessary."
    )

    add_heading(doc, "5. Required Hardware Components", 1)
    add_section_intro(
        doc,
        "The mandatory setup is only a normal laptop or desktop with a microphone and display."
    )
    add_hardware_table(doc)
    add_caption(doc, "Table 3. Required and optional hardware.")
    add_text(
        doc,
        "Development and verification used an Apple Silicon MacBook Air. The software also "
        "targets 64-bit Windows and Linux systems supported by the listed Python libraries. "
        "A dedicated GPU is unnecessary because recognition uses the small CPU-based Vosk model."
    )

    add_heading(doc, "6. Required Software", 1)
    add_software_table(doc)
    add_caption(doc, "Table 4. Software stack fixed for the prototype.")
    add_heading(doc, "6.1 Project Structure", 2)
    add_bullets(
        doc,
        [
            "voice_cursor/app.py - Tkinter interface, listener thread, session logging, and guided evaluation.",
            "voice_cursor/speech.py - secure model download and streaming Vosk recognizer.",
            "voice_cursor/commands.py - supported vocabulary, aliases, normalization, and exact matching.",
            "voice_cursor/actions.py - paused-state enforcement and PyAutoGUI/browser actions.",
            "voice_cursor/evaluation.py - trial records, accuracy/response-time summaries, and CSV export.",
            "tests/ - automated tests for commands, actions, and evaluation logic.",
        ],
    )
    add_heading(doc, "6.2 Installation and Execution", 2)
    add_steps(
        doc,
        [
            "Create and activate a Python virtual environment in the project folder.",
            "Install the project with “python -m pip install -e .”.",
            "Run “voice-cursor --dry-run” first and allow microphone access.",
            "Verify recognition, pause/resume behavior, and the guided evaluation.",
            "Grant operating-system Accessibility permission and run “voice-cursor” for live actions.",
        ],
    )

    add_heading(doc, "7. Finalized Methodology", 1)
    add_heading(doc, "7.1 Development Method", 2)
    add_steps(
        doc,
        [
            "Translate literature findings into requirements: offline privacy, fixed discoverable commands, visible feedback, manual fallback, and measurable outcomes.",
            "Capture 16 kHz mono audio and decode it with a pretrained Vosk model restricted to the supported grammar.",
            "Normalize final recognized text and map only exact phrases or aliases to typed command objects.",
            "Pass the command through a safety-state executor before any cursor, click, scroll, browser, or shortcut action.",
            "Show transcript and result feedback and write a timestamped CSV interaction log.",
            "Use dry-run and guided evaluation modes before live operating-system control.",
        ],
    )
    add_heading(doc, "7.2 Experimental Design for Human Evaluation", 2)
    add_text(
        doc,
        "Design: within-participant comparison of quiet and moderate-noise conditions. Recruit "
        "10 adult volunteers for a formative study; include varied accents where feasible and "
        "record accent/language background only with consent. Each participant completes the "
        "12 non-state commands once per condition, producing 240 command trials. Command order "
        "is randomized by the application. Actions remain disabled during measurement."
    )
    add_heading(doc, "7.3 Variables and Measures", 2)
    add_bullets(
        doc,
        [
            "Independent variable: acoustic condition (quiet vs. moderate background noise).",
            "Primary dependent variable: command accuracy = correct recognized commands / total prompts × 100.",
            "Efficiency measure: response time from prompt display to final recognized result.",
            "Error measures: substitution, unrecognized result, timeout, and repeated attempt.",
            "Subjective measures: concise SASSI-aligned ratings for accuracy, speed, cognitive demand, annoyance, likeability, and habitability.",
            "Safety measure: number of unintended actions while control is paused or evaluation mode is active; required value is zero.",
        ],
    )
    add_heading(doc, "7.4 Procedure", 2)
    add_steps(
        doc,
        [
            "Provide an information sheet, obtain consent, explain that the prototype is not a medical device, and demonstrate the pause control.",
            "Collect minimal non-identifying background information relevant to speech interaction.",
            "Allow a two-minute practice in dry-run mode with all supported commands visible.",
            "Run the randomized quiet-condition evaluation and save its CSV output.",
            "Run the moderate-noise evaluation at a documented sound level and save a second CSV.",
            "Administer the short SASSI-aligned questionnaire and ask one open-ended improvement question.",
            "Remove names from filenames, retain only derived CSV data, and delete any incidental recordings; the application itself does not store audio.",
        ],
    )
    add_heading(doc, "7.5 Analysis Plan and Acceptance Criteria", 2)
    add_text(
        doc,
        "Report mean, median, standard deviation, and range for accuracy and response time. "
        "Compare paired quiet/noise results with a Wilcoxon signed-rank test because the sample "
        "is small and normality cannot be assumed. Summarize errors by command and participant "
        "rather than reporting only an overall percentage. The prototype acceptance targets are "
        "at least 85% command accuracy in quiet conditions, median response time no greater than "
        "3 seconds, and zero actions executed while paused or in evaluation mode."
    )
    add_heading(doc, "7.6 Ethics, Privacy, and Limitations", 2)
    add_bullets(
        doc,
        [
            "Obtain informed consent and permit withdrawal without penalty.",
            "Do not recruit clinically vulnerable users for DA1 without institutional approval and accessible procedures.",
            "Do not retain raw voice recordings; store only target, recognized text, timing, and correctness.",
            "Report accent and disability limitations and avoid presenting one-speaker accuracy as universal performance.",
            "Do not automate destructive actions such as file deletion, purchases, messaging, or security settings.",
            "Treat the system as a research prototype, not a certified assistive or medical product.",
        ],
    )

    add_heading(doc, "8. Implementation Verification and Expected Outcomes", 1)
    add_heading(doc, "8.1 Completed Verification", 2)
    add_test_table(doc)
    add_caption(doc, "Table 5. Verification evidence for the implemented prototype.")
    add_callout(
        doc,
        "Interpretation of 14/14 result",
        "This result proves that the end-to-end model, grammar, and parser can recognize every "
        "supported primary phrase in clean synthesized audio. It does not prove real-world "
        "accuracy, accent fairness, disability accessibility, or usability; those require the "
        "human evaluation in Section 7.",
    )
    add_heading(doc, "8.2 Expected Outcomes", 2)
    add_bullets(
        doc,
        [
            "A reproducible low-cost desktop demonstration requiring no specialized hardware.",
            "Quantified accuracy and response-time data for quiet and noisy environments.",
            "Identification of commands that are easily confused or slow to recognize.",
            "Evidence about perceived accuracy, speed, effort, annoyance, and learnability.",
            "A clear basis for future adaptation to Indian English, personalized speech, or additional safe commands.",
        ],
    )
    add_heading(doc, "8.3 Limitations and Future Work", 2)
    add_text(
        doc,
        "Discrete directional commands are slower than continuous pointing for long-distance "
        "movement. Recognition may decline with noise, accents not represented in the model, "
        "dysarthria, low microphone quality, or fatigue. The current English grammar is small "
        "and the prototype has not been clinically validated. Future work may add adjustable "
        "step sizes, dwell or grid-based pointing, multilingual models, personalized grammars, "
        "optional text-to-speech feedback, and ethically approved studies with target users."
    )

    add_heading(doc, "9. Conclusion", 1)
    add_text(
        doc,
        "The finalized topic is feasible, literature-supported, inexpensive, and already "
        "implemented as a working prototype. Thirty DOI-verified papers show both the promise "
        "of speech as an alternative input and the importance of discoverability, error "
        "recovery, privacy, multimodal fallback, user diversity, and honest evaluation. The "
        "chosen pretrained offline approach removes the need for custom data collection while "
        "retaining genuine AI-based speech recognition. The hardware, software, and methodology "
        "are fixed, and the remaining research activity is a clearly defined human evaluation "
        "rather than unfinished system development."
    )

    add_heading(doc, "References", 1)
    add_section_intro(
        doc,
        "All 30 scholarly references include a clickable DOI link, as required."
    )
    add_references(doc)

    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_report())
